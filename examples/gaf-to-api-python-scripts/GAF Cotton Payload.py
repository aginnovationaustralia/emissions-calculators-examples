import openpyxl
import requests
import json

# =========================
# CONFIG
# =========================
XLSX_FILE = r"your path here"
API_URL   = "https://emissionscalculator-mtls.production.aiaapi.com/calculator/3.0.2/cotton"
CERT_PATH = r"your path here"
KEY_PATH  = r"your path here"

# =========================
# LOAD WORKBOOK
# =========================
wb = openpyxl.load_workbook(XLSX_FILE, data_only=True)
ws_crops = wb["Data input - crops"]
ws_veg   = wb["Data input - vegetation"]

# =========================
# SAFE VALUE HELPERS
# =========================
def f(ws, col, row, default=0.0):
    v = ws[f"{col}{row}"].value
    try:
        return float(v)
    except (TypeError, ValueError):
        return default

def s(ws, col, row, default=""):
    v = ws[f"{col}{row}"].value
    return str(v).strip() if v is not None else default

def yes_no(ws, col, row):
    """Read a 'Yes'/'No' cell into a boolean."""
    return s(ws, col, row).strip().lower() in ("yes", "y", "true", "1")

def lookup_code(ws, col, row, table, label, default):
    raw = ws[f"{col}{row}"].value
    try:
        code = int(raw)
    except (TypeError, ValueError):
        print(f"  ! {label}: could not read code {raw!r} at {col}{row} - using default {default!r}")
        return default
    if code in table:
        return table[code]
    print(f"  ! {label}: unknown code {code!r} at {col}{row} - using default {default!r}")
    return default

# =========================
# LOOKUP TABLES
# Region code -> API state string.  Source: Electricity!R2:S10
#   1 ACT | 2 NSW | 3 Tas | 4 SW WA | 5 SA | 6 Vic | 7 Qld | 8 NT | 9 NW WA
# NOTE: wa_sw / wa_nw are the still-unconfirmed feedlot enums - irrelevant
#       for this NSW farm, flagged for when a WA workbook comes through.
# =========================
STATE_LOOKUP = {
    1: "act",
    2: "nsw",
    3: "tas",
    4: "wa_sw",
    5: "sa",
    6: "vic",
    7: "qld",
    8: "nt",
    9: "wa_nw",
}

# =========================
# CROP CELL CONSTANTS  (Data input - crops, all values in column C)
# Cotton is a single-crop vertical layout.
# >>> ENTERPRISE-SPECIFIC
# =========================
CC = "C"
ROW_REGION          = 2    # region code -> STATE_LOOKUP
ROW_RAINFALL        = 5    # 'Yes'/'No'  -> rainfallAbove600
ROW_BALE_WEIGHT     = 8    # kg/bale     -> averageWeightPerBaleKg
ROW_YIELD_BALES_HA  = 9    # bales/ha    -> averageCottonYield  (NOT row 7, which is t/farm)
ROW_LINT_PER_BALE   = 10   # kg/bale     -> cottonLintPerBaleKg
ROW_SEED_PER_BALE   = 11   # kg/bale     -> cottonSeedPerBaleKg
ROW_WASTE_PER_BALE  = 12   # kg/bale     -> wastePerBaleKg
ROW_AREA_SOWN       = 13   # ha/farm     -> areaSown
ROW_NONUREA_N       = 15   # kg N/ha     -> nonUreaNitrogen
ROW_UREA            = 16   # kg Urea/ha  -> ureaApplication
ROW_UAN             = 17   # kg/ha       -> ureaAmmoniumNitrate
ROW_PHOSPHORUS      = 18   # kg P/ha     -> phosphorusApplication
ROW_POTASSIUM       = 19   # kg K/ha     -> potassiumApplication
ROW_SULFUR          = 20   # kg S/ha     -> sulfurApplication
ROW_LIME            = 21   # total t     -> limestone
ROW_LIME_FRACTION   = 22   # fraction    -> limestoneFraction
ROW_DIESEL          = 24   # litres/yr   -> dieselUse
ROW_PETROL          = 25   # litres/yr   -> petrolUse
ROW_LPG             = 26   # litres/yr   -> lpg
ROW_ELEC_USE        = 27   # kWh         -> electricityUse (farm-level)
ROW_ELEC_RENEWABLE  = 28   # 0-100%      -> electricityRenewable (farm-level)
ROW_HERBICIDE       = 29   # kg a.i.     -> herbicideUse
ROW_GLYPHOSATE      = 30   # kg a.i.     -> glyphosateOtherHerbicideUse

# =========================
# FARM-LEVEL FIELDS
# =========================
state                = lookup_code(ws_crops, CC, ROW_REGION, STATE_LOOKUP, "state", "nsw")
electricity_use      = f(ws_crops, CC, ROW_ELEC_USE)
electricity_renew    = f(ws_crops, CC, ROW_ELEC_RENEWABLE)

# =========================
# BUILD CROP
# Single cotton crop; electricity fully allocated to it.
# otherFertiliserApplication & singleSuperPhosphate are NOT captured in this
# workbook version -> zeroed (see note in delivery message).
# =========================
crops = [{
    "id":                          "",
    "state":                       state,
    "averageCottonYield":          f(ws_crops, CC, ROW_YIELD_BALES_HA),
    "areaSown":                    f(ws_crops, CC, ROW_AREA_SOWN),
    "averageWeightPerBaleKg":      f(ws_crops, CC, ROW_BALE_WEIGHT),
    "cottonLintPerBaleKg":         f(ws_crops, CC, ROW_LINT_PER_BALE),
    "cottonSeedPerBaleKg":         f(ws_crops, CC, ROW_SEED_PER_BALE),
    "wastePerBaleKg":              f(ws_crops, CC, ROW_WASTE_PER_BALE),
    "ureaApplication":             f(ws_crops, CC, ROW_UREA),
    "otherFertiliserApplication":  0.0,   # not in workbook
    "nonUreaNitrogen":             f(ws_crops, CC, ROW_NONUREA_N),
    "ureaAmmoniumNitrate":         f(ws_crops, CC, ROW_UAN),
    "phosphorusApplication":       f(ws_crops, CC, ROW_PHOSPHORUS),
    "potassiumApplication":        f(ws_crops, CC, ROW_POTASSIUM),
    "sulfurApplication":           f(ws_crops, CC, ROW_SULFUR),
    "singleSuperPhosphate":        0.0,   # not in workbook
    "rainfallAbove600":            yes_no(ws_crops, CC, ROW_RAINFALL),
    "herbicideUse":                f(ws_crops, CC, ROW_HERBICIDE),
    "glyphosateOtherHerbicideUse": f(ws_crops, CC, ROW_GLYPHOSATE),
    "electricityAllocation":       1.0,   # single crop -> full allocation
    "limestone":                   f(ws_crops, CC, ROW_LIME),
    "limestoneFraction":           f(ws_crops, CC, ROW_LIME_FRACTION),
    "dieselUse":                   f(ws_crops, CC, ROW_DIESEL),
    "petrolUse":                   f(ws_crops, CC, ROW_PETROL),
    "lpg":                         f(ws_crops, CC, ROW_LPG),
}]
num_crops = len(crops)

# =========================
# BUILD VEGETATION
# Vertical repeating blocks on 'Data input - vegetation': label 'State' in
# column C, values in column E.  Offsets from the State row:
#   +1 Region | +2 Species | +3 Soil | +4 Area | +5 Age
# Only blocks with area > 0 are included; empty ones are skipped.
# Repeating arrays must never be empty -> single zeroed default fallback.
# =========================
def build_vegetation():
    blocks = []
    for row in range(1, ws_veg.max_row + 1):
        if ws_veg[f"C{row}"].value != "State":
            continue
        area = f(ws_veg, "E", row + 4)
        if area <= 0:
            continue
        blocks.append({
            "vegetation": {
                "region":      s(ws_veg, "E", row + 1),
                "treeSpecies": s(ws_veg, "E", row + 2),
                "soil":        s(ws_veg, "E", row + 3),
                "area":        area,
                "age":         f(ws_veg, "E", row + 5),
            },
            "allocationToCrops": [0.0] * num_crops,
        })

    if not blocks:
        blocks.append({
            "vegetation": {
                "region":      "South West",
                "treeSpecies": "Mixed species (Environmental Plantings)",
                "soil":        "Loams & Clays",
                "area":        0.0,
                "age":         0.0,
            },
            "allocationToCrops": [0.0] * num_crops,
        })
    return blocks

vegetation = build_vegetation()

# =========================
# BUILD FULL PAYLOAD
# =========================
payload = {
    "id":                   "",
    "state":                state,
    "crops":                crops,
    "electricityRenewable": electricity_renew,
    "electricityUse":       electricity_use,
    "vegetation":           vegetation,
}

# =========================
# PAYLOAD PREVIEW
# =========================
print("=" * 25)
print("PAYLOAD PREVIEW")
print("=" * 25)
print(json.dumps(payload, indent=2))
print("=" * 25)

# =========================
# POST TO API
# =========================
headers = {"Content-Type": "application/json"}

response = requests.post(
    API_URL,
    json=payload,
    headers=headers,
    cert=(CERT_PATH, KEY_PATH),
    verify=True,
)

print("=" * 25)
print("API RESPONSE")
print("=" * 25)
print("Status code:", response.status_code)
print(response.text)
