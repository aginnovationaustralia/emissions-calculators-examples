import json
import warnings

import openpyxl
import requests

warnings.simplefilter("ignore")  # openpyxl DrawingML / data-validation UserWarnings

# =========================
# CONFIG
# =========================
XLSX_PATH = r"your path here"
CERT_PATH = r"your path here"
KEY_PATH = r"your path here"
API_URL = "https://emissionscalculator-mtls.production.aiaapi.com/calculator/3.0.2/pork"

# =========================
# LOAD WORKBOOK
# =========================
wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
di = wb["Data input"]
mm = wb["Manure management"]
pf = wb["Pig Feed"]
veg_ws = wb["Data input - vegetation"]


# =========================
# SAFE VALUE HELPERS
# =========================
def f(ws, coord, default=0.0):
    """Float value from a cell."""
    v = ws[coord].value
    if v is None or str(v).strip() == "":
        return default
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def s(ws, coord, default=""):
    """String value from a cell."""
    v = ws[coord].value
    if v is None:
        return default
    return str(v).strip()


def yes_no(ws, coord, default=False):
    """Yes/No dropdown -> bool."""
    v = s(ws, coord).lower()
    if v in ("yes", "y", "true"):
        return True
    if v in ("no", "n", "false"):
        return False
    return default


def normalize_enum(value, lookup, label):
    """Map a workbook label onto an API enum, warning if it isn't recognised."""
    key = str(value).strip().lower()
    for k, v in lookup.items():
        if k.lower() == key:
            return v
    print(f"WARNING: unrecognised {label} value '{value}' — sending as-is")
    return str(value).strip()


# =========================
# ENTERPRISE-SPECIFIC LOOKUPS  >>> ENTERPRISE-SPECIFIC
# =========================
# 'Data input' D4 region label -> API state enum
STATE_LOOKUP = {
    "NSW": "nsw",
    "VIC": "vic",
    "QLD": "qld",
    "SA": "sa",
    "WA": "wa",
    "SW WA": "wa",
    "NW WA": "wa",
    "TAS": "tas",
    "NT": "nt",
}

# API class key -> ('Data input' column, VS rate label on 'Manure management')
# VS rates: Manure management R6:S9 (Boars 0.40, Sows 0.46, Gilts 0.55,
# Slaughter pig herd 0.39 — used for suckers, weaners, growers, slaughter pigs)
PORK_CLASSES = {
    "sows": ("C", "S7"),
    "boars": ("D", "S6"),
    "gilts": ("E", "S8"),
    "suckers": ("F", "S9"),
    "weaners": ("G", "S9"),
    "growers": ("H", "S9"),
    "slaughterPigs": ("I", "S9"),
}

SEASONS = ["spring", "summer", "autumn", "winter"]

# 'Data input' row numbers
ROW_HEAD = {"spring": 11, "summer": 12, "autumn": 13, "winter": 14}
ROW_RAW_SOLIDS = {"spring": 39, "summer": 40, "autumn": 41, "winter": 42}
ROW_VS_PERCENT = 44
ROW_PURCHASE_HEAD = 18
ROW_PURCHASE_WEIGHT = 19
ROW_HEAD_SOLD = 23
ROW_SALE_WEIGHT = 24

# Manure system -> 'Data input' percentage row
MANURE_ROWS = {
    "uncoveredAnaerobicPond": 51,
    "coveredAnaerobicPond": 52,
    "deepLitter": 53,
    "outdoorSystems": 54,
}

# 'GWP Factors' C55:C58 — days in each season
SEASON_DAYS = {"spring": 91, "summer": 90, "autumn": 92, "winter": 92}

# Feed products: 'Data input' columns C:F, rows 29-32
FEED_COLS = ["C", "D", "E", "F"]
ROW_FEED_PRODUCT = 29
ROW_FEED_PURCHASED = 30
ROW_FEED_ADDITIONAL = 31
ROW_FEED_INTENSITY = 32

# 'Pig Feed' ration table — (product name cell, ingredient col, percentage col)
FEED_RATION_COLS = [
    ("C3", "B", "C"),
    ("F3", "E", "F"),
    ("I3", "H", "I"),
    ("L3", "K", "L"),
    ("O3", "N", "O"),
]
FEED_RATION_ROWS = range(6, 14)  # ingredient rows on 'Pig Feed'

# 'Pig Feed' ingredient label -> API ingredient key
INGREDIENT_LOOKUP = {
    "wheat": "wheat",
    "barley": "barley",
    "whey powder": "wheyPowder",
    "canola meal": "canolaMeal",
    "soybean meal": "soybeanMeal",
    "meat meal": "meatMeal",
    "blood meal": "bloodMeal",
    "fishmeal": "fishmeal",
    "tallow": "tallow",
    "wheat bran": "wheatBran",
    "beet pulp": "beetPulp",
    "millmix": "millMix",
    "mill mix": "millMix",
}

DEFAULT_VEGETATION = {
    "vegetation": {
        "region": "South West",
        "treeSpecies": "Mixed species (Environmental Plantings)",
        "soil": "Loams & Clays",
        "area": 0,
        "age": 0,
    },
    "allocatedProportion": [0],
}


# =========================
# BUILDERS
# =========================
def volatile_solids(col, vs_rate_cell, season):
    """Tonnes of volatile solids produced by one class in one season.

    Mirrors 'Data input' rows 46-49:
      manual mode  -> raw solids (rows 39-42) x volatile fraction (row 44)
      default mode -> head x VS rate (kg/hd/day) x days in season x 10^-3
    """
    if s(di, "C36").lower().startswith("enter manually"):
        raw = f(di, f"{col}{ROW_RAW_SOLIDS[season]}")
        return raw * f(di, f"{col}{ROW_VS_PERCENT}")

    head = f(di, f"{col}{ROW_HEAD[season]}")
    rate = f(mm, vs_rate_cell)
    return head * rate * SEASON_DAYS[season] * 1e-3


def build_manure(col, vs_rate_cell):
    """Tonnes of volatile solids allocated to each manure system, per season."""
    manure = {}
    for season in SEASONS:
        vs = volatile_solids(col, vs_rate_cell, season)
        entry = {
            system: vs * f(di, f"{col}{row}")
            for system, row in MANURE_ROWS.items()
        }
        entry["undefinedSystem"] = 0  # no workbook input row for undefined systems
        manure[season] = entry
    return manure


def build_purchases(col):
    """Never returns an empty array."""
    return [
        {
            "head": f(di, f"{col}{ROW_PURCHASE_HEAD}"),
            "purchaseWeight": f(di, f"{col}{ROW_PURCHASE_WEIGHT}"),
        }
    ]


def build_class(col, vs_rate_cell):
    return {
        "autumn": f(di, f"{col}{ROW_HEAD['autumn']}"),
        "winter": f(di, f"{col}{ROW_HEAD['winter']}"),
        "spring": f(di, f"{col}{ROW_HEAD['spring']}"),
        "summer": f(di, f"{col}{ROW_HEAD['summer']}"),
        "headSold": f(di, f"{col}{ROW_HEAD_SOLD}"),
        "saleWeight": f(di, f"{col}{ROW_SALE_WEIGHT}"),
        "purchases": build_purchases(col),
        "manure": build_manure(col, vs_rate_cell),
    }


def build_fertiliser():
    others = [
        {
            "otherType": s(di, "C61"),
            "otherDryland": f(di, "C60"),
            "otherIrrigated": f(di, "E60"),
        }
    ]
    return {
        "singleSuperphosphate": f(di, "C63"),
        "pastureDryland": f(di, "C58"),
        "pastureIrrigated": f(di, "E58"),
        "cropsDryland": f(di, "C59"),
        "cropsIrrigated": f(di, "E59"),
        "otherFertilisers": others,
    }


def build_ration_table():
    """Product name -> {apiIngredientKey: fraction} from the 'Pig Feed' sheet."""
    rations = {}
    for name_cell, ing_col, pct_col in FEED_RATION_COLS:
        product = s(pf, name_cell)
        if not product:
            continue
        mix = {}
        for row in FEED_RATION_ROWS:
            label = s(pf, f"{ing_col}{row}")
            if not label:
                continue
            key = normalize_enum(label, INGREDIENT_LOOKUP, "feed ingredient")
            mix[key] = f(pf, f"{pct_col}{row}")
        rations[product.lower()] = mix
    return rations


def build_feed_products():
    """Never returns an empty array — feeds with 0 tonnes purchased are skipped."""
    rations = build_ration_table()
    all_keys = sorted({k for mix in rations.values() for k in mix})

    feed_products = []
    for col in FEED_COLS:
        purchased = f(di, f"{col}{ROW_FEED_PURCHASED}")
        if purchased <= 0:
            continue

        product = s(di, f"{col}{ROW_FEED_PRODUCT}")
        mix = rations.get(product.lower())
        if mix is None:
            print(f"WARNING: no ration found on 'Pig Feed' for product '{product}'")
            mix = {}

        ingredients = {key: mix.get(key, 0) for key in all_keys}
        feed_products.append(
            {
                "feedPurchased": purchased,
                "additionalIngredients": f(di, f"{col}{ROW_FEED_ADDITIONAL}"),
                "emissionsIntensity": f(di, f"{col}{ROW_FEED_INTENSITY}"),
                "ingredients": ingredients,
            }
        )

    if not feed_products:
        feed_products.append(
            {
                "feedPurchased": 0,
                "additionalIngredients": 0,
                "emissionsIntensity": 0,
                "ingredients": {key: 0 for key in all_keys} or {"wheat": 0},
            }
        )
    return feed_products


def build_vegetation():
    """Vegetation blocks start at row 2 and repeat every 8 rows.
    Only blocks with an area greater than zero are sent; never returns an empty array."""
    vegetation = []
    row = 2
    while s(veg_ws, f"C{row}").lower() == "state":
        area = f(veg_ws, f"D{row + 4}")
        if area > 0:
            vegetation.append(
                {
                    "vegetation": {
                        "region": s(veg_ws, f"D{row + 1}"),
                        "treeSpecies": s(veg_ws, f"D{row + 2}"),
                        "soil": s(veg_ws, f"D{row + 3}"),
                        "area": area,
                        "age": f(veg_ws, f"D{row + 5}"),
                    },
                    "allocatedProportion": [1],
                }
            )
        row += 8

    if not vegetation:
        vegetation.append(json.loads(json.dumps(DEFAULT_VEGETATION)))
    return vegetation


# =========================
# BUILD PAYLOAD
# =========================
pork_entry = {
    "id": "",
    "classes": {
        api_key: build_class(col, vs_rate_cell)
        for api_key, (col, vs_rate_cell) in PORK_CLASSES.items()
    },
    "limestone": f(di, "C65"),
    "limestoneFraction": f(di, "C66"),
    "fertiliser": build_fertiliser(),
    "diesel": f(di, "C77"),
    "petrol": f(di, "C78"),
    "lpg": f(di, "C79"),
    "electricitySource": "State Grid",
    "electricityRenewable": f(di, "C76"),
    "electricityUse": f(di, "C75"),
    "herbicide": f(di, "C68"),
    "herbicideOther": f(di, "C69"),
    "beddingHayBarleyStraw": f(di, "C73"),
    "feedProducts": build_feed_products(),
}

payload = {
    "id": "",
    "state": normalize_enum(s(di, "D4"), STATE_LOOKUP, "state"),
    "rainfallAbove600": yes_no(di, "D6"),
    "pork": [pork_entry],
    "vegetation": build_vegetation(),
}

# =========================
# PAYLOAD PREVIEW
# =========================
print("=" * 25)
print("PAYLOAD PREVIEW")
print("=" * 25)
print(json.dumps(payload, indent=2))
print("=" * 25)
print(f"Farm name: {s(di, 'J1')}")
print(f"Feed products sent: {len(pork_entry['feedProducts'])}")
print(f"Vegetation entries: {len(payload['vegetation'])}")
print("=" * 25)

# =========================
# POST TO API
# =========================
headers = {"Content-Type": "application/json"}

response = requests.post(
    API_URL,
    headers=headers,
    json=payload,
    cert=(CERT_PATH, KEY_PATH),
    verify=True,
)

print("=" * 25)
print("API RESPONSE")
print("=" * 25)
print(f"Status code: {response.status_code}")
print(response.text)