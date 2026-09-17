import json
import warnings

import openpyxl
import requests

warnings.simplefilter("ignore")  # openpyxl DrawingML / data-validation UserWarnings

# =========================
# CONFIG
# =========================
XLSX_PATH = r"your path here"
CERT_PATH = r"your path here"
KEY_PATH  = r"your path here"
API_URL   = "https://emissionscalculator-mtls.production.aiaapi.com/calculator/3.0.2/poultry"

# =========================
# SHEET NAMES
# =========================
SHEET_BROILERS = "Data input - Broilers"
SHEET_LAYERS   = "Data input - Layers"
SHEET_FEED     = "Poultry Feed"
SHEET_ELEC     = "Electricity"

# The Poultry GAF workbook does not ask the Tropic of Capricorn question.
# Set this manually if the farm sits north of the Tropic of Capricorn.
NORTH_OF_TROPIC_OF_CAPRICORN = False

# -------- Farm-level: read from the BROILERS sheet as primary --------
# (These fields appear on both sheets; on a real single-farm workbook they
#  should match. The script warns if the Layers sheet disagrees.)
B_CELL_STATE_NAME   = "D9"
B_CELL_STATE_CODE   = "D10"
B_CELL_RAINFALL_600 = "D12"
L_CELL_STATE_NAME   = "E3"
L_CELL_STATE_CODE   = "E4"
L_CELL_RAINFALL_600 = "E5"

# ========================================================================
# BROILER SHEET LAYOUT
# ========================================================================
# Bird groups: 4 groups, identical block shape, 23-row stride
B_GROUP_COUNT       = 4
B_GROUP_STRIDE      = 23
B_GROUP_ROW_BIRDS   = 18
B_GROUP_ROW_STAY50  = 19
B_GROUP_ROW_LW50    = 20
B_GROUP_ROW_STAY100 = 21
B_GROUP_ROW_LW100   = 22
B_GROUP_ROW_DMI     = 23
B_GROUP_ROW_DMD     = 24
B_GROUP_ROW_CP      = 25
B_GROUP_ROW_ASH     = 26
B_GROUP_ROW_NRR     = 27
B_COL_GROWERS = "C"
B_COL_LAYERS  = "D"
B_COL_OTHER   = "E"

# Feed block (per group): 3 feed products + custom feed, same 23-row stride
B_FEED_ROW_FIRST  = 32
B_FEED_ROWS       = 3
B_FEED_ROW_CUSTOM = 36
B_FEED_COL_TONNES = "C"
B_FEED_COL_CODE   = "D"   # reference no. -> Poultry Feed lookup
B_FEED_COL_ADDL   = "E"
B_FEED_COL_EI     = "F"

# Energy & fuel
B_CELL_ELEC_USE   = "C110"
B_CELL_ELEC_RENEW = "C111"
B_CELL_DIESEL     = "C112"
B_CELL_PETROL     = "C113"
B_CELL_LPG        = "C114"
B_CELL_HAY        = "C115"
B_CELL_HERBICIDE  = "C116"
B_CELL_HERB_OTHER = "C117"

# Manure / litter
B_CELL_MANURE_ALLOCATION = "C121"
B_CELL_WASTE_DRYLOT      = "E121"
B_CELL_LITTER_RECYCLED   = "C125"
B_CELL_LITTER_FREQUENCY  = "C126"

# Purchase inventory
B_ROW_PURCHASE_HEAD   = 130
B_ROW_PURCHASE_WEIGHT = 131
B_PUR_COL_GROWERS = "D"
B_PUR_COL_LAYERS  = "E"
B_PUR_COL_OTHER   = "F"
B_CELL_FREE_RANGE = "E136"

# Sale inventory: L1 and L2 blocks -> two entries in the API "sales" array
B_SALE_BLOCKS = [
    {"row_head": 143, "row_weight": 144},   # Sale (L1)
    {"row_head": 149, "row_weight": 150},   # Sale (L2)
]
B_SALE_COL_GROWERS = "D"
B_SALE_COL_LAYERS  = "E"
B_SALE_COL_OTHER   = "F"

B_CELL_ELEC_SOURCE = "C3"   # Electricity sheet, broiler block

# ========================================================================
# LAYER SHEET LAYOUT
# ========================================================================
# Single-value inputs live in column C; per-class data in D (Layers) / E (Meat chicken layers)
L_COL_LAYERS = "D"
L_COL_MCL    = "E"

# Flock numbers (rows 9-12): Spring, Summer, Autumn, Winter
L_ROW_FLOCK_SPRING = 9
L_ROW_FLOCK_SUMMER = 10
L_ROW_FLOCK_AUTUMN = 11
L_ROW_FLOCK_WINTER = 12

# Purchase inventory (rows 16-17)
L_ROW_PURCHASE_HEAD   = 16
L_ROW_PURCHASE_WEIGHT = 17

L_CELL_FREE_RANGE = "E22"

# Eggs sale inventory (rows 26-27)
L_ROW_EGG_PRODUCED = 26
L_ROW_EGG_WEIGHT   = 27

# Feed block (rows 41-43) + custom feed (row 45)
# NB: product reference no. is in column B on this sheet (not D as on Broilers)
L_FEED_ROW_FIRST  = 41
L_FEED_ROWS       = 3
L_FEED_ROW_CUSTOM = 45
L_FEED_COL_CODE   = "B"
L_FEED_COL_TONNES = "C"
L_FEED_COL_ADDL   = "E"
L_FEED_COL_EI     = "F"

# Energy & fuel (values in column C)
L_CELL_ELEC_USE   = "C49"   # Annual Electricity Use (State Grid); row 54 is a separate total
L_CELL_ELEC_RENEW = "C50"
L_CELL_DIESEL     = "C51"
L_CELL_PETROL     = "C52"
L_CELL_LPG        = "C53"
L_CELL_HAY        = "C55"
L_CELL_HERBICIDE  = "C56"
L_CELL_HERB_OTHER = "C57"

# Manure / litter
L_CELL_MANURE_ALLOCATION = "C61"
L_CELL_WASTE_DRYLOT      = "E61"
L_CELL_LITTER_RECYCLED   = "C65"
L_CELL_LITTER_FREQUENCY  = "C66"

L_CELL_ELEC_SOURCE = "C15"  # Electricity sheet, layer block

# ========================================================================
# SHARED LOOKUPS
# ========================================================================
# >>> ENTERPRISE-SPECIFIC: state codes (Electricity sheet, cols R/S)
STATE_LOOKUP = {
    1: "act",
    2: "nsw",
    3: "tas",
    4: "wa",    # SW WA
    5: "sa",
    6: "vic",
    7: "qld",
    8: "nt",
    9: "wa",    # NW WA
}
STATE_NAME_LOOKUP = {
    "act": "act", "nsw": "nsw", "tas": "tas", "vic": "vic", "qld": "qld",
    "sa": "sa", "nt": "nt", "wa": "wa", "sw wa": "wa", "nw wa": "wa",
}

# >>> ENTERPRISE-SPECIFIC: feed product ingredient ration mix
# Poultry Feed sheet, "Ingredients +10% of ration" table (rows 72-84)
FEED_LOOKUP_ROWS = range(72, 85)
FEED_LOOKUP_COL_CODE    = "D"
FEED_LOOKUP_COL_WHEAT   = "E"
FEED_LOOKUP_COL_BARLEY  = "F"
FEED_LOOKUP_COL_SOYBEAN = "G"
FEED_LOOKUP_COL_SORGHUM = "H"
FEED_LOOKUP_COL_MILLRUN = "I"

VALID_ELECTRICITY_SOURCES = ["State Grid", "Renewable"]

# =========================
# LOAD WORKBOOK
# =========================
wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
ws_broilers = wb[SHEET_BROILERS]
ws_layers   = wb[SHEET_LAYERS]
ws_feed     = wb[SHEET_FEED]
ws_elec     = wb[SHEET_ELEC]


# =========================
# HELPERS
# =========================
def num(ws, coord, default=0.0):
    """Numeric cell value, tolerating blanks, text zeros and stray strings."""
    v = ws[coord].value
    if v is None or str(v).strip() == "":
        return default
    try:
        return float(str(v).strip())
    except ValueError:
        return default


def txt(ws, coord, default=""):
    """Trimmed text cell value."""
    v = ws[coord].value
    if v is None:
        return default
    return str(v).strip()


def yes_no(ws, coord, default=False):
    """Yes/No cell -> bool."""
    v = txt(ws, coord).lower()
    if v in ("yes", "y", "true"):
        return True
    if v in ("no", "n", "false"):
        return False
    return default


def normalize_enum(value, valid_values, default):
    """Match a workbook string to an API enum, warning on unknown values."""
    for valid in valid_values:
        if str(value).strip().lower() == valid.lower():
            return valid
    if str(value).strip():
        print(f"  WARNING: unrecognised value '{value}' - defaulting to '{default}'")
    return default


def decode_state(code, name):
    """Decode a numeric state code, falling back to the state name."""
    if int(code) in STATE_LOOKUP:
        return STATE_LOOKUP[int(code)]
    if name.lower() in STATE_NAME_LOOKUP:
        return STATE_NAME_LOOKUP[name.lower()]
    print(f"  WARNING: could not decode state (code={code}, name='{name}') - defaulting to 'nsw'")
    return "nsw"


def build_feed_lookup():
    """Map feed product reference no. -> ingredient ration mix."""
    table = {}
    for r in FEED_LOOKUP_ROWS:
        code = ws_feed[f"{FEED_LOOKUP_COL_CODE}{r}"].value
        if code is None or str(code).strip() == "":
            continue
        table[int(code)] = {
            "wheat":   num(ws_feed, f"{FEED_LOOKUP_COL_WHEAT}{r}"),
            "barley":  num(ws_feed, f"{FEED_LOOKUP_COL_BARLEY}{r}"),
            "sorghum": num(ws_feed, f"{FEED_LOOKUP_COL_SORGHUM}{r}"),
            "soybean": num(ws_feed, f"{FEED_LOOKUP_COL_SOYBEAN}{r}"),
            "millrun": num(ws_feed, f"{FEED_LOOKUP_COL_MILLRUN}{r}"),
        }
    return table


FEED_INGREDIENTS = build_feed_lookup()


def lookup_code(code):
    """Decode a feed product code into its ingredient mix."""
    if code is None or str(code).strip() == "":
        return {"wheat": 0, "barley": 0, "sorghum": 0, "soybean": 0, "millrun": 0}
    try:
        key = int(float(code))
    except ValueError:
        key = None
    if key not in FEED_INGREDIENTS:
        print(f"  WARNING: unknown feed product code '{code}' - sending zeroed ingredients")
        return {"wheat": 0, "barley": 0, "sorghum": 0, "soybean": 0, "millrun": 0}
    return dict(FEED_INGREDIENTS[key])


def default_feed():
    """Never-empty-array fallback for feed."""
    return {
        "ingredients": {"wheat": 0, "barley": 0, "sorghum": 0, "soybean": 0, "millrun": 0},
        "feedPurchased": 0,
        "additionalIngredient": 0,
        "emissionIntensity": 0,
    }


# ========================================================================
# BROILER BUILDERS
# ========================================================================
def b_build_meat_class(col, offset):
    return {
        "birds":                  num(ws_broilers, f"{col}{B_GROUP_ROW_BIRDS + offset}"),
        "averageStayLength50":    num(ws_broilers, f"{col}{B_GROUP_ROW_STAY50 + offset}"),
        "liveweight50":           num(ws_broilers, f"{col}{B_GROUP_ROW_LW50 + offset}"),
        "averageStayLength100":   num(ws_broilers, f"{col}{B_GROUP_ROW_STAY100 + offset}"),
        "liveweight100":          num(ws_broilers, f"{col}{B_GROUP_ROW_LW100 + offset}"),
        "dryMatterIntake":        num(ws_broilers, f"{col}{B_GROUP_ROW_DMI + offset}"),
        "dryMatterDigestibility": num(ws_broilers, f"{col}{B_GROUP_ROW_DMD + offset}"),
        "crudeProtein":           num(ws_broilers, f"{col}{B_GROUP_ROW_CP + offset}"),
        "manureAsh":              num(ws_broilers, f"{col}{B_GROUP_ROW_ASH + offset}"),
        "nitrogenRetentionRate":  num(ws_broilers, f"{col}{B_GROUP_ROW_NRR + offset}"),
    }


def b_build_feed(offset):
    feed = []
    for i in range(B_FEED_ROWS):
        r = B_FEED_ROW_FIRST + offset + i
        feed.append({
            "ingredients":          lookup_code(ws_broilers[f"{B_FEED_COL_CODE}{r}"].value),
            "feedPurchased":        num(ws_broilers, f"{B_FEED_COL_TONNES}{r}"),
            "additionalIngredient": num(ws_broilers, f"{B_FEED_COL_ADDL}{r}"),
            "emissionIntensity":    num(ws_broilers, f"{B_FEED_COL_EI}{r}"),
        })
    if not feed:
        feed = [default_feed()]
    return feed


def b_build_group(offset):
    r_custom = B_FEED_ROW_CUSTOM + offset
    return {
        "meatChickenGrowers": b_build_meat_class(B_COL_GROWERS, offset),
        "meatChickenLayers":  b_build_meat_class(B_COL_LAYERS, offset),
        "meatOther":          b_build_meat_class(B_COL_OTHER, offset),
        "feed": b_build_feed(offset),
        "customFeedPurchased":         num(ws_broilers, f"{B_FEED_COL_TONNES}{r_custom}"),
        "customFeedEmissionIntensity": num(ws_broilers, f"{B_FEED_COL_EI}{r_custom}"),
    }


def b_default_group():
    zero_class = {
        "birds": 0, "averageStayLength50": 0, "liveweight50": 0,
        "averageStayLength100": 0, "liveweight100": 0, "dryMatterIntake": 0,
        "dryMatterDigestibility": 0, "crudeProtein": 0, "manureAsh": 0,
        "nitrogenRetentionRate": 0,
    }
    return {
        "meatChickenGrowers": dict(zero_class),
        "meatChickenLayers":  dict(zero_class),
        "meatOther":          dict(zero_class),
        "feed": [default_feed()],
        "customFeedPurchased": 0,
        "customFeedEmissionIntensity": 0,
    }


def b_build_purchases(col):
    return {
        "head":           num(ws_broilers, f"{col}{B_ROW_PURCHASE_HEAD}"),
        "purchaseWeight": num(ws_broilers, f"{col}{B_ROW_PURCHASE_WEIGHT}"),
    }


def b_build_sale(block):
    rh, rw = block["row_head"], block["row_weight"]
    return {
        "meatChickenGrowersSales": {
            "head":       num(ws_broilers, f"{B_SALE_COL_GROWERS}{rh}"),
            "saleWeight": num(ws_broilers, f"{B_SALE_COL_GROWERS}{rw}"),
        },
        "meatChickenLayers": {
            "head":       num(ws_broilers, f"{B_SALE_COL_LAYERS}{rh}"),
            "saleWeight": num(ws_broilers, f"{B_SALE_COL_LAYERS}{rw}"),
        },
        "meatOther": {
            "head":       num(ws_broilers, f"{B_SALE_COL_OTHER}{rh}"),
            "saleWeight": num(ws_broilers, f"{B_SALE_COL_OTHER}{rw}"),
        },
    }


def b_default_sale():
    return {
        "meatChickenGrowersSales": {"head": 0, "saleWeight": 0},
        "meatChickenLayers":       {"head": 0, "saleWeight": 0},
        "meatOther":               {"head": 0, "saleWeight": 0},
    }


def build_broiler():
    """Single broiler (chicken meat) enterprise from the Broilers sheet."""
    groups = []
    for g in range(B_GROUP_COUNT):
        offset = g * B_GROUP_STRIDE
        birds = sum(
            num(ws_broilers, f"{col}{B_GROUP_ROW_BIRDS + offset}")
            for col in (B_COL_GROWERS, B_COL_LAYERS, B_COL_OTHER)
        )
        if birds > 0:
            groups.append(b_build_group(offset))
    if not groups:
        groups = [b_default_group()]

    sales = [b_build_sale(block) for block in B_SALE_BLOCKS]
    if not sales:
        sales = [b_default_sale()]

    return {
        "id": "",
        "groups": groups,
        "diesel":  num(ws_broilers, B_CELL_DIESEL),
        "petrol":  num(ws_broilers, B_CELL_PETROL),
        "lpg":     num(ws_broilers, B_CELL_LPG),
        "electricitySource":    normalize_enum(txt(ws_elec, B_CELL_ELEC_SOURCE),
                                               VALID_ELECTRICITY_SOURCES, "State Grid"),
        "electricityRenewable": num(ws_broilers, B_CELL_ELEC_RENEW),
        "electricityUse":       num(ws_broilers, B_CELL_ELEC_USE),
        "hay":            num(ws_broilers, B_CELL_HAY),
        "herbicide":      num(ws_broilers, B_CELL_HERBICIDE),
        "herbicideOther": num(ws_broilers, B_CELL_HERB_OTHER),
        "manureWasteAllocation":       num(ws_broilers, B_CELL_MANURE_ALLOCATION),
        "wasteHandledDrylotOrStorage": num(ws_broilers, B_CELL_WASTE_DRYLOT),
        "litterRecycled":         num(ws_broilers, B_CELL_LITTER_RECYCLED),
        "litterRecycleFrequency": num(ws_broilers, B_CELL_LITTER_FREQUENCY),
        "purchasedFreeRange":     num(ws_broilers, B_CELL_FREE_RANGE),
        "meatChickenGrowersPurchases": b_build_purchases(B_PUR_COL_GROWERS),
        "meatChickenLayersPurchases":  b_build_purchases(B_PUR_COL_LAYERS),
        "meatOtherPurchases":          b_build_purchases(B_PUR_COL_OTHER),
        "sales": sales,
    }


# ========================================================================
# LAYER BUILDERS
# ========================================================================
def l_build_seasons(col):
    return {
        "autumn": num(ws_layers, f"{col}{L_ROW_FLOCK_AUTUMN}"),
        "winter": num(ws_layers, f"{col}{L_ROW_FLOCK_WINTER}"),
        "spring": num(ws_layers, f"{col}{L_ROW_FLOCK_SPRING}"),
        "summer": num(ws_layers, f"{col}{L_ROW_FLOCK_SUMMER}"),
    }


def l_build_feed():
    feed = []
    for i in range(L_FEED_ROWS):
        r = L_FEED_ROW_FIRST + i
        feed.append({
            "ingredients":          lookup_code(ws_layers[f"{L_FEED_COL_CODE}{r}"].value),
            "feedPurchased":        num(ws_layers, f"{L_FEED_COL_TONNES}{r}"),
            "additionalIngredient": num(ws_layers, f"{L_FEED_COL_ADDL}{r}"),
            "emissionIntensity":    num(ws_layers, f"{L_FEED_COL_EI}{r}"),
        })
    if not feed:
        feed = [default_feed()]
    return feed


def l_build_purchases(col):
    return {
        "head":           num(ws_layers, f"{col}{L_ROW_PURCHASE_HEAD}"),
        "purchaseWeight": num(ws_layers, f"{col}{L_ROW_PURCHASE_WEIGHT}"),
    }


def l_build_egg_sale(col):
    return {
        "eggsProduced":  num(ws_layers, f"{col}{L_ROW_EGG_PRODUCED}"),
        "averageWeight": num(ws_layers, f"{col}{L_ROW_EGG_WEIGHT}"),
    }


def build_layer():
    """Single layer (egg) enterprise from the Layers sheet."""
    return {
        "id": "",
        "layers":            l_build_seasons(L_COL_LAYERS),
        "meatChickenLayers": l_build_seasons(L_COL_MCL),
        "feed": l_build_feed(),
        "purchasedFreeRange": num(ws_layers, L_CELL_FREE_RANGE),
        "diesel":  num(ws_layers, L_CELL_DIESEL),
        "petrol":  num(ws_layers, L_CELL_PETROL),
        "lpg":     num(ws_layers, L_CELL_LPG),
        "electricitySource":    normalize_enum(txt(ws_elec, L_CELL_ELEC_SOURCE),
                                               VALID_ELECTRICITY_SOURCES, "State Grid"),
        "electricityRenewable": num(ws_layers, L_CELL_ELEC_RENEW),
        "electricityUse":       num(ws_layers, L_CELL_ELEC_USE),
        "hay":            num(ws_layers, L_CELL_HAY),
        "herbicide":      num(ws_layers, L_CELL_HERBICIDE),
        "herbicideOther": num(ws_layers, L_CELL_HERB_OTHER),
        "manureWasteAllocation":       num(ws_layers, L_CELL_MANURE_ALLOCATION),
        "wasteHandledDrylotOrStorage": num(ws_layers, L_CELL_WASTE_DRYLOT),
        "litterRecycled":         num(ws_layers, L_CELL_LITTER_RECYCLED),
        "litterRecycleFrequency": num(ws_layers, L_CELL_LITTER_FREQUENCY),
        "meatChickenLayersPurchases": l_build_purchases(L_COL_MCL),
        "layersPurchases":            l_build_purchases(L_COL_LAYERS),
        "customFeedPurchased":         num(ws_layers, f"{L_FEED_COL_TONNES}{L_FEED_ROW_CUSTOM}"),
        "customFeedEmissionIntensity": num(ws_layers, f"{L_FEED_COL_EI}{L_FEED_ROW_CUSTOM}"),
        "meatChickenLayersEggSale": l_build_egg_sale(L_COL_MCL),
        "layersEggSale":            l_build_egg_sale(L_COL_LAYERS),
    }


# ========================================================================
# VEGETATION
# ========================================================================
def default_vegetation(n_broilers, n_layers):
    """Never-empty-array fallback for vegetation."""
    return {
        "vegetation": {
            "region": "South West",
            "treeSpecies": "Mixed species (Environmental Plantings)",
            "soil": "Loams & Clays",
            "area": 0,
            "age": 0,
        },
        "broilersProportion": [0] * max(n_broilers, 1),
        "layersProportion":   [0] * max(n_layers, 1),
    }


# =========================
# FARM-LEVEL FIELDS (Broilers sheet primary, warn on Layers mismatch)
# =========================
state_broiler = decode_state(num(ws_broilers, B_CELL_STATE_CODE),
                             txt(ws_broilers, B_CELL_STATE_NAME))
state_layer   = decode_state(num(ws_layers, L_CELL_STATE_CODE),
                             txt(ws_layers, L_CELL_STATE_NAME))
rain_broiler  = yes_no(ws_broilers, B_CELL_RAINFALL_600)
rain_layer    = yes_no(ws_layers, L_CELL_RAINFALL_600)

if state_broiler != state_layer:
    print(f"  WARNING: state differs between sheets "
          f"(Broilers='{state_broiler}', Layers='{state_layer}'). "
          f"Using Broilers value - confirm which is correct.")
if rain_broiler != rain_layer:
    print(f"  WARNING: rainfall-above-600 differs between sheets "
          f"(Broilers={rain_broiler}, Layers={rain_layer}). "
          f"Using Broilers value - confirm which is correct.")

# =========================
# BUILD PAYLOAD
# =========================
broilers = [build_broiler()]
layers   = [build_layer()]
vegetation = [default_vegetation(len(broilers), len(layers))]

payload = {
    "state": state_broiler,
    "northOfTropicOfCapricorn": NORTH_OF_TROPIC_OF_CAPRICORN,
    "rainfallAbove600": rain_broiler,
    "broilers": broilers,
    "layers": layers,
    "vegetation": vegetation,
}

# =========================
# PAYLOAD PREVIEW
# =========================
print("=" * 25)
print("PAYLOAD PREVIEW")
print("=" * 25)
print(json.dumps(payload, indent=2))
print("=" * 25)
print(f"Broiler enterprises: {len(broilers)}")
for idx, ent in enumerate(broilers):
    print(f"  broilers[{idx}]: {len(ent['groups'])} group(s), {len(ent['sales'])} sales entry(s)")
print(f"Layer enterprises  : {len(layers)}")
for idx, ent in enumerate(layers):
    print(f"  layers[{idx}]: {len(ent['feed'])} feed entry(s)")
print(f"Vegetation entries : {len(vegetation)}")
print("=" * 25)

# =========================
# POST TO API
# =========================
headers = {"Content-Type": "application/json"}

response = requests.post(
    API_URL,
    headers=headers,
    json=payload,
    cert=(CERT_PATH, KEY_PATH),
    verify=True,
)

print("=" * 25)
print("API RESPONSE")
print("=" * 25)
print(f"Status code: {response.status_code}")
print(response.text)