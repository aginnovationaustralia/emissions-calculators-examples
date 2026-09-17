import openpyxl
import requests
import json

# =========================
# CONFIG
# =========================
# >>> ENTERPRISE-SPECIFIC: path to the completed G-GAF workbook
XLSX_FILE = r"your path here"
API_URL   = "https://emissionscalculator-mtls.production.aiaapi.com/calculator/3.0.2/rice"
CERT_FILE = r"your path here"
KEY_FILE  = r"your path here"

# >>> ENTERPRISE-SPECIFIC: sheet names inside the G-GAF workbook
CROPS_SHEET      = "Data input - crops"
VEGETATION_SHEET = "Data input - vegetation"

# =========================
# LOAD WORKBOOK
# =========================
wb = openpyxl.load_workbook(XLSX_FILE, data_only=True)
ws_crops = wb[CROPS_SHEET]
ws_veg   = wb[VEGETATION_SHEET]

# =========================
# SAFE VALUE HELPERS
# =========================
def f(value, default=0.0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default

def s(value, default=""):
    if value is None:
        return default
    return str(value).strip().strip('"')

# =========================
# >>> ENTERPRISE-SPECIFIC: LOOKUP TABLES
# The Rice G-GAF workbook stores the farm region as a numeric code,
# resolved via the lookup table at 'Leaching and runoff'!L6:M14.
# Unlike the grains workbook, crop type is fixed (Rice) and the
# water regime fields are entered directly as dropdown strings,
# so no crop-type or production-system lookups are needed.
# The production system code (C4) and orange zone code (C5) drive
# internal workbook calculations only - the rice API schema has no
# corresponding fields.
# =========================

# NOTE: WA is split into two rainfall sub-regions in the workbook
# (SW WA / NW WA) but both resolve to the same API state code "wa".
STATE_LOOKUP = {
    1: "act",
    2: "nsw",
    3: "tas",
    4: "wa",   # SW WA
    5: "sa",
    6: "vic",
    7: "qld",
    8: "nt",
    9: "wa",   # NW WA
}

def lookup_code(code_value, table, label):
    try:
        code = int(f(code_value))
    except (TypeError, ValueError):
        code = None
    if code not in table:
        print(f"WARNING: unrecognised {label} code: {code_value!r}")
        return ""
    return table[code]

# =========================
# >>> ENTERPRISE-SPECIFIC: ENUM NORMALISATION
# The workbook's pre-season flooding dropdown strings (see
# 'Methane Rice'!A83:A86) omit spaces around the < / > signs,
# but the API enum values include them, e.g. workbook
# 'Non flooded pre-season >180 days' must be sent as
# 'Non flooded pre-season > 180 days' (confirmed against the
# working EAP Rice 3.0.2 flat CSV). Unmapped values pass
# through unchanged with a warning.
# =========================
FLOODING_PERIOD_NORMALISE = {
    "Non flooded pre-season <180 days":  "Non flooded pre-season < 180 days",
    "Non flooded pre-season >180 days":  "Non flooded pre-season > 180 days",
    "Flooded pre-season >30 days":       "Flooded pre-season > 30 days",
    "Non-flooded pre-season >365 days":  "Non-flooded pre-season > 365 days",
}

def normalise_flooding_period(value):
    raw = s(value)
    if raw in FLOODING_PERIOD_NORMALISE:
        return FLOODING_PERIOD_NORMALISE[raw]
    print(f"WARNING: pre-season flooding period not in normalisation map, "
          f"sending as-is: {raw!r}")
    return raw

# =========================
# FARM-LEVEL AND CROP FIELDS
# The rice workbook uses a single-crop VERTICAL layout: every value
# sits in column C of "Data input - crops", one field per row
# (unlike the grains workbook, where crops run across columns).
# >>> ENTERPRISE-SPECIFIC: row numbers within column C
# =========================
VALUE_COL = 3  # column C

ROW_REGION_CODE          = 2
ROW_AVG_RICE_YIELD       = 6
ROW_AREA_SOWN            = 7
ROW_GROWING_SEASON_DAYS  = 8
ROW_WATER_REGIME_TYPE    = 9
ROW_WATER_REGIME_SUBTYPE = 10
ROW_PRESEASON_FLOODING   = 11
ROW_NON_UREA_NITROGEN    = 13
ROW_UREA_APPLICATION     = 14
ROW_UREA_AMMONIUM_NITRATE = 15
ROW_PHOSPHORUS           = 16
ROW_POTASSIUM            = 17
ROW_SULFUR               = 18
ROW_LIMESTONE            = 19
ROW_LIMESTONE_FRACTION   = 20
ROW_FRACTION_CROP_BURNT  = 21
ROW_DIESEL               = 22
ROW_PETROL               = 23
ROW_LPG                  = 24
ROW_ELECTRICITY_USE      = 25
ROW_ELECTRICITY_RENEW    = 26
ROW_ELECTRICITY_ALLOC    = 27
ROW_HERBICIDE_USE        = 28
ROW_GLYPHOSATE_OTHER     = 29

def crop_cell(row):
    return ws_crops.cell(row=row, column=VALUE_COL).value

farm_state = lookup_code(crop_cell(ROW_REGION_CODE), STATE_LOOKUP, "region")

# =========================
# BUILD CROPS LIST
# Single rice crop per workbook.
# =========================
crops = [{
    "id":                          "",
    "state":                       farm_state,
    "averageRiceYield":            f(crop_cell(ROW_AVG_RICE_YIELD)),
    "areaSown":                    f(crop_cell(ROW_AREA_SOWN)),
    "growingSeasonDays":           f(crop_cell(ROW_GROWING_SEASON_DAYS)),
    "waterRegimeType":             s(crop_cell(ROW_WATER_REGIME_TYPE)),
    "waterRegimeSubType":          s(crop_cell(ROW_WATER_REGIME_SUBTYPE)),
    "ricePreseasonFloodingPeriod": normalise_flooding_period(crop_cell(ROW_PRESEASON_FLOODING)),
    "ureaApplication":             f(crop_cell(ROW_UREA_APPLICATION)),
    "nonUreaNitrogen":             f(crop_cell(ROW_NON_UREA_NITROGEN)),
    "ureaAmmoniumNitrate":         f(crop_cell(ROW_UREA_AMMONIUM_NITRATE)),
    "phosphorusApplication":       f(crop_cell(ROW_PHOSPHORUS)),
    "potassiumApplication":        f(crop_cell(ROW_POTASSIUM)),
    "sulfurApplication":           f(crop_cell(ROW_SULFUR)),
    "fractionOfAnnualCropBurnt":   f(crop_cell(ROW_FRACTION_CROP_BURNT)),
    "herbicideUse":                f(crop_cell(ROW_HERBICIDE_USE)),
    "glyphosateOtherHerbicideUse": f(crop_cell(ROW_GLYPHOSATE_OTHER)),
    "electricityAllocation":       f(crop_cell(ROW_ELECTRICITY_ALLOC)),
    "limestone":                   f(crop_cell(ROW_LIMESTONE)),
    "limestoneFraction":           f(crop_cell(ROW_LIMESTONE_FRACTION)),
    "dieselUse":                   f(crop_cell(ROW_DIESEL)),
    "petrolUse":                   f(crop_cell(ROW_PETROL)),
    "lpg":                         f(crop_cell(ROW_LPG)),
}]

num_crop_slots = len(crops)

# =========================
# BUILD VEGETATION LIST
# The rice vegetation sheet holds four fixed placeholder blocks
# (labels in column B, values in column D):
#   State / Region / Species of Tree / Soil Type / Area of Trees / Age of Trees
# There are NO "Allocation to crop" rows in this workbook (unlike
# grains), so any block with a real planted area is allocated 100%
# to the single rice crop.
# Blocks are detected by scanning column B for the label "State"
# rather than assuming fixed row spacing, so this is robust to
# layout drift in future workbook versions.
# Zero-area placeholder blocks are skipped; if none remain, a single
# zeroed default entry is sent (repeating arrays must never be empty).
# =========================
LABEL_COL = 2  # column B holds the row labels within a vegetation block
VEG_VALUE_COL = 4  # column D holds the corresponding values

vegetation = []
row = 1
max_row = ws_veg.max_row
while row <= max_row:
    label = ws_veg.cell(row=row, column=LABEL_COL).value
    if label != "State":
        row += 1
        continue

    region_row  = row + 1
    species_row = row + 2
    soil_row    = row + 3
    area_row    = row + 4
    age_row     = row + 5

    area = f(ws_veg.cell(row=area_row, column=VEG_VALUE_COL).value)
    if area > 0:
        vegetation.append({
            "vegetation": {
                "age":         f(ws_veg.cell(row=age_row, column=VEG_VALUE_COL).value),
                "area":        area,
                "region":      s(ws_veg.cell(row=region_row, column=VEG_VALUE_COL).value),
                "soil":        s(ws_veg.cell(row=soil_row, column=VEG_VALUE_COL).value),
                "treeSpecies": s(ws_veg.cell(row=species_row, column=VEG_VALUE_COL).value),
            },
            "allocationToCrops": [1.0] * num_crop_slots,
        })

    row = age_row + 1  # jump past this block

# Repeating arrays must never be empty - default to a single zeroed entry
if not vegetation:
    vegetation.append({
        "vegetation": {
            "age":         0.0,
            "area":        0.0,
            "region":      "South West",
            "soil":        "Loams & Clays",
            "treeSpecies": "Mixed species (Environmental Plantings)",
        },
        "allocationToCrops": [0.0] * num_crop_slots,
    })

# =========================
# BUILD FULL PAYLOAD
# =========================
payload = {
    "state":               farm_state,
    "electricityUse":      f(crop_cell(ROW_ELECTRICITY_USE)),
    "electricityRenewable":f(crop_cell(ROW_ELECTRICITY_RENEW)),
    "crops":               crops,
    "vegetation":          vegetation,
}

# =========================
# PREVIEW
# =========================
print("=== PAYLOAD PREVIEW ===")
print(json.dumps(payload, indent=2))

# =========================
# SEND REQUEST
# =========================
headers = {"Content-Type": "application/json"}

response = requests.post(
    API_URL,
    json=payload,
    headers=headers,
    cert=(CERT_FILE, KEY_FILE),
    verify=True
)

print("\n=== RESPONSE ===")
print("Status:", response.status_code)
print(response.text)