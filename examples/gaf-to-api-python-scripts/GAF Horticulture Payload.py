import openpyxl
import requests
import json

# =========================
# CONFIG
# =========================
XLSX_FILE = r"your path here"
API_URL   = "https://emissionscalculator-mtls.production.aiaapi.com/calculator/3.0.2/horticulture"
CERT_PATH = r"your path here"
KEY_PATH  = r"your path here"

# =========================
# LOAD WORKBOOK
# =========================
wb = openpyxl.load_workbook(XLSX_FILE, data_only=True)
ws_crops = wb["Data input - crops"]
ws_veg   = wb["Data input - vegetation"]

# =========================
# SAFE VALUE HELPERS
# =========================
def f(ws, col, row, default=0.0):
    v = ws[f"{col}{row}"].value
    try:
        return float(v)
    except (TypeError, ValueError):
        return default

def s(ws, col, row, default=""):
    v = ws[f"{col}{row}"].value
    return str(v).strip() if v is not None else default

def lookup_code(ws, col, row, table, label, default):
    raw = ws[f"{col}{row}"].value
    try:
        code = int(raw)
    except (TypeError, ValueError):
        print(f"  ! {label}: could not read code {raw!r} at {col}{row} - using default {default!r}")
        return default
    if code in table:
        return table[code]
    print(f"  ! {label}: unrecognised code {code} at {col}{row} - using default {default!r}")
    return default

# =========================
# LOOKUP TABLES  (decode numeric workbook codes)
# Region codes:  Electricity!R2:S10
# Crop type codes: Crop Residues!B41:C46
# =========================
STATE_LOOKUP = {
    1: "act",
    2: "nsw",
    3: "tas",
    4: "wa_sw",   # >>> ENTERPRISE-SPECIFIC: SW WA enum unconfirmed - verify against a live 200
    5: "sa",
    6: "vic",
    7: "qld",
    8: "nt",
    9: "wa_nw",   # >>> ENTERPRISE-SPECIFIC: NW WA enum unconfirmed - verify against a live 200
}

CROP_TYPE_LOOKUP = {
    1: "Pulses",
    2: "Tuber and Roots",
    3: "Peanuts",
    4: "Hops",
    5: "Perennial Hort",
    6: "Annual Hort",
}

# =========================
# FIXED CELL POSITIONS  (Data input - crops)
# Crops sit horizontally: columns C, D, E, F  ->  crop slots 0-3
# >>> ENTERPRISE-SPECIFIC: re-check these rows if the workbook version changes
# =========================
CROP_COLS = ["C", "D", "E", "F"]

ROW_REGION_CODE           = 2    # col C only (farm-level)
ROW_CROP_TYPE             = 4
ROW_RAINFALL              = 5    # 1 = Yes (>600mm / drains) -> true, else false
ROW_YIELD                 = 6
ROW_AREA                  = 7
ROW_NON_UREA_N            = 8
ROW_PHOSPHORUS            = 9
ROW_POTASSIUM             = 10
ROW_SULFUR                = 11
ROW_UREA                  = 12
ROW_UAN                   = 13
ROW_LIME                  = 14
ROW_LIME_FRACTION         = 15
ROW_FRACTION_BURNT        = 16
ROW_DIESEL                = 17
ROW_PETROL                = 18
ROW_LPG                   = 19
ROW_ELECTRICITY_USE       = 20   # col C only (farm-level)
ROW_ELECTRICITY_RENEWABLE = 21   # col C only (farm-level)
ROW_ALLOCATION            = 22   # per-crop electricity allocation
ROW_REFRIGERANT           = 23   # per-crop refrigerant name
ROW_RECHARGE              = 24   # per-crop annual recharge (chargeSize)
ROW_HERBICIDE             = 25
ROW_GLYPHOSATE            = 26

# =========================
# FARM-LEVEL FIELDS
# =========================
state = lookup_code(ws_crops, "C", ROW_REGION_CODE, STATE_LOOKUP, "region", "sa")
electricity_use       = f(ws_crops, "C", ROW_ELECTRICITY_USE)
electricity_renewable = f(ws_crops, "C", ROW_ELECTRICITY_RENEWABLE)

# =========================
# BUILD CROPS  (all four fixed slots, matching the Grains GAF convention)
# =========================
def build_crop(col):
    name   = s(ws_crops, col, ROW_REFRIGERANT, "HFC-23")
    charge = f(ws_crops, col, ROW_RECHARGE)
    # Refrigerants array must never be empty - one entry always (name is set per slot)
    refrigerants = [{"refrigerant": name, "chargeSize": charge}]

    return {
        "id":                          "",
        "type":                        lookup_code(ws_crops, col, ROW_CROP_TYPE, CROP_TYPE_LOOKUP, "crop type", "Perennial Hort"),
        "averageYield":                f(ws_crops, col, ROW_YIELD),
        "areaSown":                    f(ws_crops, col, ROW_AREA),
        "ureaApplication":             f(ws_crops, col, ROW_UREA),
        "nonUreaNitrogen":             f(ws_crops, col, ROW_NON_UREA_N),
        "ureaAmmoniumNitrate":         f(ws_crops, col, ROW_UAN),
        "phosphorusApplication":       f(ws_crops, col, ROW_PHOSPHORUS),
        "potassiumApplication":        f(ws_crops, col, ROW_POTASSIUM),
        "sulfurApplication":           f(ws_crops, col, ROW_SULFUR),
        "rainfallAbove600":            int(f(ws_crops, col, ROW_RAINFALL, 2)) == 1,
        "fractionOfAnnualCropBurnt":   f(ws_crops, col, ROW_FRACTION_BURNT),
        "herbicideUse":                f(ws_crops, col, ROW_HERBICIDE),
        "glyphosateOtherHerbicideUse": f(ws_crops, col, ROW_GLYPHOSATE),
        "electricityAllocation":       f(ws_crops, col, ROW_ALLOCATION),
        "limestone":                   f(ws_crops, col, ROW_LIME),
        "limestoneFraction":           f(ws_crops, col, ROW_LIME_FRACTION),
        "dieselUse":                   f(ws_crops, col, ROW_DIESEL),
        "petrolUse":                   f(ws_crops, col, ROW_PETROL),
        "lpg":                         f(ws_crops, col, ROW_LPG),
        "refrigerants":                refrigerants,
    }

crops = [build_crop(col) for col in CROP_COLS]
num_crops = len(crops)

# =========================
# BUILD VEGETATION  (vertical repeating blocks; scan col B for "State")
# Layout per block:  State / Region / Species / Soil / Area / Age
#                    then 4 "Allocation to crop" rows (one per crop slot)
# Zero-area blocks skipped; single zeroed default used as fallback.
# =========================
def build_vegetation():
    blocks = []
    for row in range(1, ws_veg.max_row + 1):
        if ws_veg[f"B{row}"].value != "State":
            continue
        area = f(ws_veg, "D", row + 4)
        if area <= 0:
            continue
        allocs = [f(ws_veg, "D", row + 7 + k) for k in range(num_crops)]
        blocks.append({
            "id": "",
            "vegetation": {
                "region":      s(ws_veg, "D", row + 1),
                "treeSpecies": s(ws_veg, "D", row + 2),
                "soil":        s(ws_veg, "D", row + 3),
                "area":        area,
                "age":         f(ws_veg, "D", row + 5),
            },
            "allocationToCrops": allocs,
        })

    # Repeating arrays must never be empty - default to a single zeroed entry
    if not blocks:
        blocks.append({
            "id": "",
            "vegetation": {
                "region":      "South West",
                "treeSpecies": "Mixed species (Environmental Plantings)",
                "soil":        "Loams & Clays",
                "area":        0.0,
                "age":         0.0,
            },
            "allocationToCrops": [0.0] * num_crops,
        })
    return blocks

vegetation = build_vegetation()

# =========================
# BUILD FULL PAYLOAD
# =========================
payload = {
    "id":                   "",
    "state":                state,
    "crops":                crops,
    "electricityRenewable": electricity_renewable,
    "electricityUse":       electricity_use,
    "vegetation":           vegetation,
}

# =========================
# PAYLOAD PREVIEW
# =========================
print("=" * 25)
print("PAYLOAD PREVIEW")
print("=" * 25)
print(json.dumps(payload, indent=2))
print("=" * 25)

# =========================
# POST TO API
# =========================
headers = {"Content-Type": "application/json"}

response = requests.post(
    API_URL,
    json=payload,
    headers=headers,
    cert=(CERT_PATH, KEY_PATH),
    verify=True,
)

print("=" * 25)
print("API RESPONSE")
print("=" * 25)
print("Status code:", response.status_code)
print(response.text)
