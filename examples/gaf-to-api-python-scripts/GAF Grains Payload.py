import openpyxl
import requests
import json

# =========================
# CONFIG
# =========================
# >>> ENTERPRISE-SPECIFIC: path to the completed G-GAF workbook
XLSX_FILE = r"your path here"
API_URL   = "https://emissionscalculator-mtls.production.aiaapi.com/calculator/calculator/3.0.0/grains"
CERT_FILE = r"your path here"
KEY_FILE  = r"your path here"

# >>> ENTERPRISE-SPECIFIC: sheet names inside the G-GAF workbook
CROPS_SHEET      = "Data input - crops"
VEGETATION_SHEET = "Data input - vegetation"

# =========================
# LOAD WORKBOOK
# =========================
wb = openpyxl.load_workbook(XLSX_FILE, data_only=True)
ws_crops = wb[CROPS_SHEET]
ws_veg   = wb[VEGETATION_SHEET]

# =========================
# SAFE VALUE HELPERS
# =========================
def f(value, default=0.0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default

def s(value, default=""):
    if value is None:
        return default
    return str(value).strip().strip('"')

def yes_no(value):
    return str(value).strip().lower() == "yes"

# =========================
# >>> ENTERPRISE-SPECIFIC: LOOKUP TABLES
# The G-GAF workbook stores crop type, production system, and farm
# region as numeric codes. These are resolved via hidden lookup tables
# elsewhere in the workbook:
#   - Crop type        -> 'Crop Residues'!B68:C85
#   - Production system -> 'Leaching and runoff'!B18:C20
#   - Farm region/state -> 'Electricity'!C2:D10 (same table repeated on
#                          'Leaching and runoff'!B6:C14)
# =========================
CROP_TYPE_LOOKUP = {
    1: "Wheat",
    2: "Barley",
    3: "Maize",
    4: "Oats",
    5: "Sorghum",
    6: "Triticale",
    7: "Other Cereals",
    8: "Pulses",
    9: "Tuber and Roots",
    10: "Peanuts",
    11: "Hops",
    12: "Oilseeds",
    13: "Forage Crops",
    14: "Lucerne",
    15: "Other legume",
    16: "Annual grass",
    17: "Grass clover mixture",
    18: "Perennial pasture",
}

PRODUCTION_SYSTEM_LOOKUP = {
    1: "Non-irrigated crop",
    2: "Irrigated crop",
    3: "Sugar cane",
}

# NOTE: WA is split into two rainfall sub-regions in the workbook
# (SW WA / NW WA) but both resolve to the same API state code "wa".
STATE_LOOKUP = {
    1: "act",
    2: "nsw",
    3: "tas",
    4: "wa",   # SW WA
    5: "sa",
    6: "vic",
    7: "qld",
    8: "nt",
    9: "wa",   # NW WA
}

def lookup_code(code_value, table, label):
    try:
        code = int(f(code_value))
    except (TypeError, ValueError):
        code = None
    if code not in table:
        print(f"WARNING: unrecognised {label} code: {code_value!r}")
        return ""
    return table[code]

# =========================
# FARM-LEVEL FIELDS
# =========================
# >>> ENTERPRISE-SPECIFIC: fixed row/column positions in "Data input - crops"
REGION_CODE_CELL       = (2, 3)   # row 2, col C
ELECTRICITY_USE_CELL   = (21, 3)  # row 21, col C
ELECTRICITY_RENEW_CELL = (22, 3)  # row 22, col C

farm_state = lookup_code(ws_crops.cell(*REGION_CODE_CELL).value, STATE_LOOKUP, "region")

# =========================
# BUILD CROPS LIST
# Crop slots run across columns C, D, E, F, G... in "Data input - crops".
# We loop until we hit a blank column (row 3 = crop type code).
# >>> ENTERPRISE-SPECIFIC: row numbers for each field within a crop column
# =========================
ROW_CROP_TYPE            = 3
ROW_PRODUCTION_SYSTEM    = 4
ROW_RAINFALL_ABOVE_600   = 5
ROW_AVG_GRAIN_YIELD      = 7
ROW_AREA_SOWN            = 8
ROW_NON_UREA_NITROGEN    = 9
ROW_PHOSPHORUS           = 10
ROW_POTASSIUM            = 11
ROW_SULFUR               = 12
ROW_UREA_APPLICATION     = 13
ROW_UREA_AMMONIUM_NITRATE = 14
ROW_LIMESTONE            = 15
ROW_LIMESTONE_FRACTION   = 16
ROW_FRACTION_CROP_BURNT  = 17
ROW_DIESEL               = 18
ROW_PETROL               = 19
ROW_LPG                  = 20
ROW_ELECTRICITY_ALLOC    = 23
ROW_HERBICIDE_USE        = 24
ROW_GLYPHOSATE_OTHER     = 25

FIRST_CROP_COLUMN = 3  # column C

crops = []
col = FIRST_CROP_COLUMN
while ws_crops.cell(row=ROW_CROP_TYPE, column=col).value is not None:
    crop_type_code = ws_crops.cell(row=ROW_CROP_TYPE, column=col).value
    prod_system_code = ws_crops.cell(row=ROW_PRODUCTION_SYSTEM, column=col).value

    crops.append({
        "id":                          "",
        "type":                        lookup_code(crop_type_code, CROP_TYPE_LOOKUP, "crop type"),
        "productionSystem":            lookup_code(prod_system_code, PRODUCTION_SYSTEM_LOOKUP, "production system"),
        "state":                       farm_state,
        "areaSown":                    f(ws_crops.cell(row=ROW_AREA_SOWN, column=col).value),
        "averageGrainYield":           f(ws_crops.cell(row=ROW_AVG_GRAIN_YIELD, column=col).value),
        "nonUreaNitrogen":             f(ws_crops.cell(row=ROW_NON_UREA_NITROGEN, column=col).value),
        "ureaApplication":             f(ws_crops.cell(row=ROW_UREA_APPLICATION, column=col).value),
        "ureaAmmoniumNitrate":         f(ws_crops.cell(row=ROW_UREA_AMMONIUM_NITRATE, column=col).value),
        "phosphorusApplication":       f(ws_crops.cell(row=ROW_PHOSPHORUS, column=col).value),
        "potassiumApplication":        f(ws_crops.cell(row=ROW_POTASSIUM, column=col).value),
        "sulfurApplication":           f(ws_crops.cell(row=ROW_SULFUR, column=col).value),
        "rainfallAbove600":            yes_no(ws_crops.cell(row=ROW_RAINFALL_ABOVE_600, column=col).value),
        "fractionOfAnnualCropBurnt":   f(ws_crops.cell(row=ROW_FRACTION_CROP_BURNT, column=col).value),
        "herbicideUse":                f(ws_crops.cell(row=ROW_HERBICIDE_USE, column=col).value),
        "glyphosateOtherHerbicideUse": f(ws_crops.cell(row=ROW_GLYPHOSATE_OTHER, column=col).value),
        "electricityAllocation":       f(ws_crops.cell(row=ROW_ELECTRICITY_ALLOC, column=col).value),
        "limestoneFraction":           f(ws_crops.cell(row=ROW_LIMESTONE_FRACTION, column=col).value),
        "limestone":                   f(ws_crops.cell(row=ROW_LIMESTONE, column=col).value),
        "dieselUse":                   f(ws_crops.cell(row=ROW_DIESEL, column=col).value),
        "petrolUse":                   f(ws_crops.cell(row=ROW_PETROL, column=col).value),
        "lpg":                         f(ws_crops.cell(row=ROW_LPG, column=col).value),
    })
    col += 1

num_crop_slots = len(crops)

# =========================
# BUILD VEGETATION LIST
# The vegetation sheet repeats a block for each vegetation entry:
#   State / Region / Species of Tree / Soil Type / Area of Trees / Age of Trees
#   then one "Allocation to crop" row per crop slot, in the same order
#   as the crop columns above (C, D, E, F, G...).
# Blocks are detected by scanning column C for the label "State" rather
# than assuming a fixed row spacing, so this is robust to layout drift
# in future workbook versions.
# =========================
LABEL_COL = 3  # column C holds the row labels within a vegetation block
VALUE_COL = 5  # column E holds the corresponding values
ALLOC_VALUE_COL = 5  # column E holds the allocation fraction

vegetation = []
row = 1
max_row = ws_veg.max_row
while row <= max_row:
    label = ws_veg.cell(row=row, column=LABEL_COL).value
    if label != "State":
        row += 1
        continue

    state_row   = row
    region_row  = row + 1
    species_row = row + 2
    soil_row    = row + 3
    area_row    = row + 4
    age_row     = row + 5
    alloc_start_row = row + 7  # one blank row separates Age from the allocations

    allocations = []
    for k in range(num_crop_slots):
        allocations.append(f(ws_veg.cell(row=alloc_start_row + k, column=ALLOC_VALUE_COL).value))

    vegetation.append({
        "vegetation": {
            "age":         f(ws_veg.cell(row=age_row, column=VALUE_COL).value),
            "area":        f(ws_veg.cell(row=area_row, column=VALUE_COL).value),
            "region":      s(ws_veg.cell(row=region_row, column=VALUE_COL).value),
            "soil":        s(ws_veg.cell(row=soil_row, column=VALUE_COL).value),
            "treeSpecies": s(ws_veg.cell(row=species_row, column=VALUE_COL).value),
        },
        "allocationToCrops": allocations,
    })

    row = alloc_start_row + num_crop_slots  # jump past this block

# =========================
# BUILD FULL PAYLOAD
# =========================
payload = {
    "id":                    "",
    "state":                 farm_state,
    "electricityUse":        f(ws_crops.cell(*ELECTRICITY_USE_CELL).value),
    "electricityRenewable":  f(ws_crops.cell(*ELECTRICITY_RENEW_CELL).value),
    "crops":                 crops,
    "vegetation":             vegetation,
}

# =========================
# PREVIEW
# =========================
print("=== PAYLOAD PREVIEW ===")
print(json.dumps(payload, indent=2))

# =========================
# SEND REQUEST
# =========================
headers = {"Content-Type": "application/json"}

response = requests.post(
    API_URL,
    json=payload,
    headers=headers,
    cert=(CERT_FILE, KEY_FILE),
    verify=True
)

print("\n=== RESPONSE ===")
print("Status:", response.status_code)
print(response.text)