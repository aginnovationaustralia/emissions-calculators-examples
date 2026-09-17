import warnings
import json
import requests
import openpyxl

# =========================
# CONFIG
# =========================
XLSX_FILE = r"your path here"
API_URL   = "https://emissionscalculator-mtls.production.aiaapi.com/calculator/3.0.2/sheep"
CERT_FILE = r"your path here"
KEY_FILE  = r"your path here"

# openpyxl warns about DrawingML / data-validation extensions it cannot parse. Harmless.
warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

# =========================
# LOAD WORKBOOK
# =========================
wb = openpyxl.load_workbook(XLSX_FILE, data_only=True)

SHEET_SHEEP = " Data input - sheep"          # NOTE: leading space is part of the real sheet name
SHEET_BEEF  = "Data input - beef"
SHEET_VEG   = "Data input - vegetation"

ws_sheep = wb[SHEET_SHEEP]
ws_beef  = wb[SHEET_BEEF]
ws_veg   = wb[SHEET_VEG]

# =========================
# SAFE VALUE HELPERS
# =========================
def f(ws, cell, default=0):
    """Float value from a worksheet cell, with fallback default."""
    try:
        v = ws[cell].value
        if v is None or v == "":
            return default
        return float(v)
    except (ValueError, TypeError):
        return default

def s(ws, cell, default=""):
    """String value from a worksheet cell, with fallback default."""
    v = ws[cell].value
    if v is None:
        return default
    v = str(v).strip()
    return v if v else default

def is_set(ws, cell):
    """True if the cell holds a non-zero, non-blank value."""
    v = ws[cell].value
    if v is None or v == "":
        return False
    try:
        return float(v) != 0
    except (ValueError, TypeError):
        return True

def yes_no(ws, cell, default=False):
    """Workbook 'Yes'/'No' dropdown -> bool."""
    v = str(ws[cell].value).strip().lower()
    if v in ("yes", "y", "true"):
        return True
    if v in ("no", "n", "false"):
        return False
    return default

def normalize_enum(value, lookup, label, default):
    """Map a workbook label onto its API enum string, warning on anything unrecognised."""
    if value in lookup:
        return lookup[value]
    print(f"WARNING: unknown {label} value {value!r} - falling back to {default!r}")
    return default

# =========================
# LOOKUPS  >>> ENTERPRISE-SPECIFIC
# =========================
# Workbook region label (sheep sheet E4) -> API `state` enum.
# NOTE: only "SW WA" -> "wa_sw" is confirmed against a validated EAP submission.
# The rest follow the same convention and should be checked the first time
# a workbook from another state is run through this script.
STATE_LOOKUP = {
    "ACT":   "act",
    "NSW":   "nsw",
    "Tas":   "tas",
    "SW WA": "wa_sw",
    "SA":    "sa",
    "Vic":   "vic",
    "Qld":   "qld",
    "NW WA": "wa_nw",
    "NT":    "nt",
}

# Workbook electricity source label -> API `electricitySource` enum.
ELECTRICITY_LOOKUP = {
    "State Grid (Default)": "State Grid",
    "State Grid":           "State Grid",
}

# =========================
# SHEEP SHEET LAYOUT  >>> ENTERPRISE-SPECIFIC
# =========================
# The sheep sheet is laid out vertically: each animal class is a COLUMN, each
# measure is a fixed ROW. Classes run across columns D..M.
#
# The workbook carries 10 classes; the API expects 14. The seven breeder classes
# map one-to-one. The workbook's three trade columns are mapped onto the API trade
# classes that share their standard reference weight (see 'Nitrous Oxide_MMS - sheep'
# R15:AB25 - Trade lambs/hoggets SRW 66 = lambs, Trade wethers SRW 72 = wethers,
# Trade ewes SRW 60 = ewes), so the emissions maths is unchanged:
#
#     Trade lambs and hoggets (K) -> tradeWetherLambs
#     Trade wethers           (L) -> tradeWethers
#     Trade ewes              (M) -> tradeOtherEwes
#
# The four remaining API trade classes have no workbook column and are sent zeroed.
# In THIS workbook all three trade columns are empty, so the choice has no effect
# on the result - but revisit it if a workbook ever carries real trade stock.
# Declared in the API schema's own class order.
CLASS_COLUMNS = {
    "rams":                    "D",
    "tradeRams":               None,  # no workbook column - sent zeroed
    "wethers":                 "E",
    "tradeWethers":            "L",   # workbook: Trade wethers
    "maidenBreedingEwes":      "F",
    "tradeMaidenBreedingEwes": None,  # no workbook column - sent zeroed
    "breedingEwes":            "G",
    "tradeBreedingEwes":       None,  # no workbook column - sent zeroed
    "otherEwes":               "H",
    "tradeOtherEwes":          "M",   # workbook: Trade ewes
    "eweLambs":                "I",
    "tradeEweLambs":           None,  # no workbook column - sent zeroed
    "wetherLambs":             "J",
    "tradeWetherLambs":        "K",   # workbook: Trade lambs and hoggets
}

# Seasonal measures: row number per season.
SEASON_ROWS = {
    "head":                   {"spring": 9,  "summer": 10, "autumn": 11, "winter": 12},
    "liveweight":             {"spring": 15, "summer": 16, "autumn": 17, "winter": 18},
    "liveweightGain":         {"spring": 21, "summer": 22, "autumn": 23, "winter": 24},
    "crudeProtein":           {"spring": 73, "summer": 74, "autumn": 75, "winter": 76},
    "dryMatterDigestibility": {"spring": 79, "summer": 80, "autumn": 81, "winter": 82},
}

# Whole-of-year measures: single row each.
CLASS_ROWS = {
    "purchaseHead":   34,
    "purchaseWeight": 35,
    "headSold":       42,
    "saleWeight":     43,
    "headShorn":      51,
    "woolShorn":      52,
    "cleanWoolYield": 55,
}

# The SB-GAF seasonal workbook uses the CP/DMD method, so feed availability
# (sheet rows 27-30) is not an API input - the API-aligned EAP export omits it
# entirely. Sent as 0 throughout, matching the validated EAP submission.
FEED_AVAILABILITY = 0

# =========================
# BUILDER FUNCTIONS
# =========================
def build_season(col, season):
    if col is None:
        return {
            "head": 0,
            "liveweight": 0,
            "liveweightGain": 0,
            "crudeProtein": 0,
            "dryMatterDigestibility": 0,
            "feedAvailability": FEED_AVAILABILITY,
        }
    return {
        "head": f(ws_sheep, f"{col}{SEASON_ROWS['head'][season]}"),
        "liveweight": f(ws_sheep, f"{col}{SEASON_ROWS['liveweight'][season]}"),
        "liveweightGain": f(ws_sheep, f"{col}{SEASON_ROWS['liveweightGain'][season]}"),
        "crudeProtein": f(ws_sheep, f"{col}{SEASON_ROWS['crudeProtein'][season]}"),
        "dryMatterDigestibility": f(ws_sheep, f"{col}{SEASON_ROWS['dryMatterDigestibility'][season]}"),
        "feedAvailability": FEED_AVAILABILITY,
    }

def build_purchases(col):
    # The workbook allows a single purchase line per class (rows 34/35).
    if col is not None and is_set(ws_sheep, f"{col}{CLASS_ROWS['purchaseHead']}"):
        return [{
            "head": f(ws_sheep, f"{col}{CLASS_ROWS['purchaseHead']}"),
            "purchaseWeight": f(ws_sheep, f"{col}{CLASS_ROWS['purchaseWeight']}"),
        }]
    return [{"head": 0, "purchaseWeight": 0}]  # never send an empty array

def build_class(col):
    return {
        "autumn": build_season(col, "autumn"),
        "winter": build_season(col, "winter"),
        "spring": build_season(col, "spring"),
        "summer": build_season(col, "summer"),
        "headShorn": f(ws_sheep, f"{col}{CLASS_ROWS['headShorn']}") if col else 0,
        "woolShorn": f(ws_sheep, f"{col}{CLASS_ROWS['woolShorn']}") if col else 0,
        "cleanWoolYield": f(ws_sheep, f"{col}{CLASS_ROWS['cleanWoolYield']}") if col else 0,
        "headSold": f(ws_sheep, f"{col}{CLASS_ROWS['headSold']}") if col else 0,
        "saleWeight": f(ws_sheep, f"{col}{CLASS_ROWS['saleWeight']}") if col else 0,
        "purchases": build_purchases(col),
    }

def build_other_fertilisers():
    # Workbook row 98: C = product name, D = dryland tonnes, F = irrigated tonnes.
    return [{
        "otherType": s(ws_sheep, "C98", "Monoammonium phosphate (MAP)"),
        "otherDryland": f(ws_sheep, "D98"),
        "otherIrrigated": f(ws_sheep, "F98"),
    }]

def build_sheep():
    # The workbook holds exactly one sheep enterprise (one data input sheet).
    classes = {}
    for api_name, col in CLASS_COLUMNS.items():
        classes[api_name] = build_class(col)

    return {
        "id": "",
        "classes": classes,
        "limestone": f(ws_sheep, "D101"),
        "limestoneFraction": f(ws_sheep, "D102"),
        "fertiliser": {
            "singleSuperphosphate": f(ws_sheep, "D100"),
            "pastureDryland": f(ws_sheep, "D96"),
            "pastureIrrigated": f(ws_sheep, "F96"),
            "cropsDryland": f(ws_sheep, "D97"),
            "cropsIrrigated": f(ws_sheep, "F97"),
            "otherFertilisers": build_other_fertilisers(),
        },
        "diesel": f(ws_sheep, "D107"),
        "petrol": f(ws_sheep, "D108"),
        "lpg": f(ws_sheep, "D109"),
        # Mineral supplementation rows 86-88: column D is the % urea of the product,
        # column F is the tonnes of product. Do not swap these.
        "mineralSupplementation": {
            "mineralBlock": f(ws_sheep, "F86"),
            "mineralBlockUrea": f(ws_sheep, "D86"),
            "weanerBlock": f(ws_sheep, "F87"),
            "weanerBlockUrea": f(ws_sheep, "D87"),
            "drySeasonMix": f(ws_sheep, "F88"),
            "drySeasonMixUrea": f(ws_sheep, "D88"),
        },
        "electricitySource": normalize_enum(
            s(ws_sheep, "C105"), ELECTRICITY_LOOKUP, "electricity source", "State Grid"
        ),
        "electricityRenewable": f(ws_sheep, "D106"),
        "electricityUse": f(ws_sheep, "D110"),
        "grainFeed": f(ws_sheep, "D111"),
        "hayFeed": f(ws_sheep, "D112"),
        "herbicide": f(ws_sheep, "D113"),
        "herbicideOther": f(ws_sheep, "D114"),
        # Row 39: D = % of purchased sheep that are Merino, E = Cross-bred.
        "merinoPercent": f(ws_sheep, "D39"),
        # Rows 61-64: proportion of ewes lambing in each season (column G).
        "ewesLambing": {
            "spring": f(ws_sheep, "G61"),
            "summer": f(ws_sheep, "G62"),
            "autumn": f(ws_sheep, "G63"),
            "winter": f(ws_sheep, "G64"),
        },
        # Rows 67-70: seasonal lambing marking rates (column G).
        "seasonalLambing": {
            "autumn": f(ws_sheep, "G69"),
            "winter": f(ws_sheep, "G70"),
            "spring": f(ws_sheep, "G67"),
            "summer": f(ws_sheep, "G68"),
        },
    }

# Vegetation blocks are stacked vertically on the vegetation sheet: four blocks,
# each 10 rows apart, values in column D. The workbook's region / species / soil
# vocabularies are the same ones the API accepts (see the ' Trees' sheet lookup
# tables), so these strings pass straight through.
VEG_BLOCK_START_ROWS = [2, 12, 22, 32]

def build_vegetation(start_row):
    return {
        "vegetation": {
            "region": s(ws_veg, f"D{start_row + 1}", "South West"),
            "treeSpecies": s(ws_veg, f"D{start_row + 2}", "Mixed species (Environmental Plantings)"),
            "soil": s(ws_veg, f"D{start_row + 3}", "Loams & Clays"),
            "area": f(ws_veg, f"D{start_row + 4}"),
            "age": f(ws_veg, f"D{start_row + 5}"),
        },
        # One proportion per sheep enterprise, and this workbook has exactly one.
        "sheepProportion": [f(ws_veg, f"D{start_row + 8}")],
    }

# =========================
# BUILD PAYLOAD
# =========================
sheep = [build_sheep()]

# Only send vegetation blocks that actually carry trees. A block with no area and
# no age sequesters nothing, so sending it adds enum risk for no benefit.
vegetation = []
for start_row in VEG_BLOCK_START_ROWS:
    if is_set(ws_veg, f"D{start_row + 4}") or is_set(ws_veg, f"D{start_row + 5}"):
        vegetation.append(build_vegetation(start_row))

if not vegetation:  # never send an empty array
    vegetation = [{
        "vegetation": {
            "region": "South West",
            "treeSpecies": "Mixed species (Environmental Plantings)",
            "soil": "Loams & Clays",
            "area": 0,
            "age": 0,
        },
        "sheepProportion": [0],
    }]

payload = {
    # Sheep sheet E4 holds the region label; F4 holds its numeric code.
    "state": normalize_enum(s(ws_sheep, "E4"), STATE_LOOKUP, "state", "wa_sw"),
    # The Tropic of Capricorn question is only asked once, on the BEEF sheet (M4),
    # and applies to the whole property - the sheep sheet does not repeat it.
    "northOfTropicOfCapricorn": yes_no(ws_beef, "M4"),
    "rainfallAbove600": yes_no(ws_sheep, "E6"),
    "sheep": sheep,
    "vegetation": vegetation,
}

# =========================
# PAYLOAD PREVIEW
# =========================
print("=== PAYLOAD PREVIEW ===")
print(json.dumps(payload, indent=2))

# =========================
# SEND TO API
# =========================
headers = {"Content-Type": "application/json"}

response = requests.post(
    API_URL,
    json=payload,
    headers=headers,
    cert=(CERT_FILE, KEY_FILE),
    verify=True
)

print("\n=== RESPONSE ===")
print("Status:", response.status_code)
print(response.text)
